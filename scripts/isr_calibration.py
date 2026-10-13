"""ISR/SCR 校准检查（选项 A 的可本地运行部分，无需 GPU）。

目的：用同一验证器(syllable_slots_v3.parse)在人类藏文文本上算 ISR/SCR，
确认验证器不会把人类文本误判成"非法"，从而支撑故事中心点
（通用模型 ISR 0.269 ≫ 人类基线 ≈0.00384 是真实相对信号，而非验证器偏差）。

注意：G1 json 只存聚合指标、不含生成文本，模型侧 ISR 不能本地重算
（需在 4090 上用同验证器重跑生成）。本脚本只校准"人类锚点"。
"""
import sys, json, os
WS = r"D:\workbuddy工作目录\2026-09-14-22-24-36"
sys.path.insert(0, WS)
from syllable_slots_v3 import split_syllables, strip_particle, parse

def isr_scr(text):
    total = inv = closed = 0
    for syl in split_syllables(text):
        if not syl.strip():
            continue
        core, part = strip_particle(syl)
        total += 1
        if core == "":
            closed += 1          # 纯粘写助词，合法且闭合
            continue
        slots, status, _ = parse(core)
        if status == "OK":
            if slots.get("vow") or slots.get("suf"):
                closed += 1
        else:
            inv += 1
    if total == 0:
        return 0.0, 0.0, 0, 0, 0
    return inv / total, closed / total, inv, closed, total

def corpus_isr(texts, label):
    TI = TC = TT = 0
    kept = 0
    for t in texts:
        if not t or not t.strip():
            continue
        i, s, inv, closed, tot = isr_scr(t)
        TI += inv; TC += closed; TT += tot; kept += 1
    isr = TI / TT if TT else 0.0
    scr = TC / TT if TT else 0.0
    print(f"[{label}] docs={kept} syll_total={TT}  ISR={isr:.5f}  SCR={scr:.5f}")
    return {"label": label, "docs": kept, "syll_total": TT,
            "isr": round(isr, 5), "scr": round(scr, 5)}

report = []

# 1) BodSUM-6000 正文（人类撰写藏文新闻，大规模干净语料）
bod = []
for f in ["data/task3/bodsum_train.jsonl", "data/task3/bodsum_test.jsonl"]:
    with open(os.path.join(WS, f), encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                d = json.loads(line)
                bod.append(d.get("text", ""))
report.append(corpus_isr(bod, "BodSUM-6000 Text (human news)"))

# 2) Ti-SUM body（人类撰写藏文新闻，聚类）
tis = []
with open(os.path.join(WS, "tisum_parsed.jsonl"), encoding="utf-8") as fh:
    for line in fh:
        line = line.strip()
        if line:
            d = json.loads(line)
            tis.append(d.get("body", ""))
report.append(corpus_isr(tis, "Ti-SUM body (human news)"))

# 3) human_eval_200.xlsx 参考列（人类续写参考，与 G1 基线同域）
try:
    import openpyxl
    cand = [os.path.join(WS, "human_eval_200.xlsx"),
            r"C:\Users\机械革命\Desktop\human_eval_200.xlsx"]
    p = next((c for c in cand if os.path.exists(c)), None)
    if p:
        wb = openpyxl.load_workbook(p)
        refs = []
        for ws in wb.worksheets:
            hdr = [(c.value or "") for c in ws[1]]
            ci = next((idx for idx, h in enumerate(hdr) if "参考" in str(h)), None)
            if ci is None:
                continue
            for r in range(2, ws.max_row + 1):
                v = ws.cell(row=r, column=ci + 1).value
                if v and str(v).strip():
                    refs.append(str(v))
        print(f"[xlsx] collected {len(refs)} 参考 from {os.path.basename(p)}")
        report.append(corpus_isr(refs, "human_eval 参考 (human continuation refs)"))
    else:
        print("[xlsx] not found, skip")
except Exception as e:
    print(f"[xlsx] skip: {e}")

# 已知 G1 模型侧（服务器同验证器报出，本地不能重算，仅列出对照）
report.append({"label": "G1 Llama-3.1-8B-Instruct gen (server, same validator)",
               "isr": 0.269, "scr": 0.0918, "note": "reported, not recomputed locally"})
report.append({"label": "G1 TiLamb-7B-LoRA gen (server, same validator)",
               "isr": 0.0043, "scr": 0.8842, "note": "reported, not recomputed locally"})
report.append({"label": "human reference baseline (ledger §203)",
               "isr": 0.00384, "scr": 0.8821, "note": "established server-side"})

with open(os.path.join(WS, "data/task3/isr_calibration.json"), "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=2)
print("\nWrote data/task3/isr_calibration.json")
